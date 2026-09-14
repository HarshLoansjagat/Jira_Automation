import os
import shutil
import time
import uuid
from html import escape
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, HTMLResponse
from sqlalchemy.orm import Session

from app.database import get_db
from app.config import settings
from app.models.report import Report
from app.models.snapshot import Sprint, Snapshot, IssueSnapshot
from app.calculations.qa import QACalculationEngine
from app.services.report_service import ReportService

router = APIRouter(prefix="/api/reports", tags=["reports"])


def _snapshot_payload(snapshot: Snapshot) -> dict:
    qa_rows = [
        {
            "Issue key": issue.issue_key,
            "Parent": issue.parent_key,
            "Summary": issue.summary,
            "Status": issue.status,
            "Issue Type": issue.issue_type,
            "Assignee": issue.assignee,
            "Story Points": issue.story_points,
            "QA Owner": issue.qa_owner,
            "QA Story Points": issue.qa_sp,
        }
        for issue in snapshot.issues
    ]
    qa_result = QACalculationEngine().calculate(qa_rows, list(qa_rows[0]) if qa_rows else [])
    return {
        "id": snapshot.id,
        "sprint_id": snapshot.sprint.id,
        "sprint": snapshot.sprint.name,
        "snapshot_date": snapshot.snapshot_date,
        "snapshot_type": snapshot.snapshot_type,
        "captured_at": snapshot.captured_at,
        "overall_scope": snapshot.overall_scope,
        "overall_completed": snapshot.overall_completed,
        "overall_remaining": snapshot.overall_remaining,
        "overall_completion": snapshot.overall_completion,
        "day1_fixed_scope": snapshot.day1_fixed_scope,
        "sprint_start": snapshot.sprint.start_date,
        "sprint_end": snapshot.sprint.end_date,
        "developers": [
            {
                "name": metric.developer_name,
                "assigned_sp": metric.assigned_sp,
                "completed_sp": metric.completed_sp,
                "remaining_sp": metric.remaining_sp,
                "completion_pct": metric.completion_pct,
            }
            for metric in snapshot.developer_metrics
        ],
        "qa": {
            "total_sp": qa_result.total_sp,
            "completed_sp": qa_result.completed_sp,
            "members": [
                {
                    "name": metric.name,
                    "assigned_sp": metric.assigned_sp,
                    "completed_sp": metric.completed_sp,
                    "remaining_sp": metric.remaining_sp,
                    "completion_pct": metric.completion_pct,
                    "task_count": metric.task_count,
                    "completed_task_count": metric.completed_task_count,
                }
                for metric in qa_result.stats.values()
            ],
        },
        "tasks": [
            {
                "issue_key": task.issue_key,
                "parent_key": task.parent_key,
                "summary": task.summary,
                "status": task.status,
                "issue_type": task.issue_type,
                "owner": task.owner,
                "story_points": task.story_points,
                "completed_sp": task.completed_sp,
                "remaining_sp": task.remaining_sp,
            }
            for task in qa_result.tasks
        ],
    }


def _delete_file_with_retry(file_path: str, retries: int = 5, delay_seconds: float = 0.25) -> None:
    """Retry file deletion to handle brief Windows Excel lock periods."""
    normalized_path = os.path.normpath(file_path)
    if not os.path.isabs(normalized_path):
        normalized_path = os.path.abspath(normalized_path)

    last_error: Optional[Exception] = None
    for attempt in range(retries):
        try:
            if os.path.exists(normalized_path):
                os.remove(normalized_path)
            return
        except FileNotFoundError:
            return
        except PermissionError as exc:
            last_error = exc
            if attempt < retries - 1:
                time.sleep(delay_seconds * (attempt + 1))
                continue
            raise

    if last_error is not None:
        raise last_error


@router.post("/upload")
async def upload_and_generate(
    eod_file: Optional[UploadFile] = File(None),
    morning_file: Optional[UploadFile] = File(None),
    snapshot_type: str = Form("EOD"),          # "MORNING" | "EOD"
    report_date: Optional[str] = Form(None),
    day1_fixed_scope: Optional[float] = Form(None),
    sprint_start: Optional[str] = Form(None),
    sprint_end: Optional[str] = Form(None),
    db: Session = Depends(get_db),
):
    """
    Upload CSVs and generate the sprint report.

    Scenario 1 — MORNING only : snapshot_type=MORNING, morning_file=<file>
    Scenario 2 — EOD only     : snapshot_type=EOD,     eod_file=<file>
    Scenario 3 — Both         : snapshot_type=EOD,     eod_file=<file>, morning_file=<file>
    """
    os.makedirs(settings.UPLOAD_STORAGE_PATH, exist_ok=True)

    # Validate: at least one file must be supplied
    morning_provided = morning_file and morning_file.filename
    eod_provided = eod_file and eod_file.filename

    if not morning_provided and not eod_provided:
        raise HTTPException(status_code=422, detail="At least one CSV file (Morning or EOD) is required.")

    # Scenario 1: morning-only — swap so we process morning as the primary file
    if snapshot_type == "MORNING":
        if not morning_provided:
            raise HTTPException(status_code=422, detail="Morning CSV is required for MORNING snapshot type.")
        primary_file = morning_file
        secondary_file = None
        original_filename = morning_file.filename
    else:
        if not eod_provided:
            raise HTTPException(status_code=422, detail="EOD CSV is required for EOD snapshot type.")
        primary_file = eod_file
        secondary_file = morning_file if morning_provided else None
        original_filename = eod_file.filename

    # Save primary file
    primary_ext = os.path.splitext(primary_file.filename)[1] or ".csv"
    # Support .csv, .tsv, .txt
    if primary_ext.lower() not in (".csv", ".tsv", ".txt"):
        primary_ext = ".csv"
    primary_path = os.path.join(
        settings.UPLOAD_STORAGE_PATH, f"{uuid.uuid4().hex}{primary_ext}"
    )
    with open(primary_path, "wb") as f:
        shutil.copyfileobj(primary_file.file, f)

    # Save secondary (morning) file if provided
    secondary_path = None
    if secondary_file:
        sec_ext = os.path.splitext(secondary_file.filename)[1] or ".csv"
        if sec_ext.lower() not in (".csv", ".tsv", ".txt"):
            sec_ext = ".csv"
        secondary_path = os.path.join(
            settings.UPLOAD_STORAGE_PATH, f"{uuid.uuid4().hex}{sec_ext}"
        )
        with open(secondary_path, "wb") as f:
            shutil.copyfileobj(secondary_file.file, f)

    try:
        service = ReportService(db)
        result = service.process_upload(
            file_path=primary_path,
            snapshot_type=snapshot_type,
            morning_file_path=secondary_path,   # only set for Scenario 3
            report_date=report_date,
            day1_fixed_scope=day1_fixed_scope,
            sprint_start=sprint_start,
            sprint_end=sprint_end,
            original_filename=original_filename,
        )
        return result
    except Exception as exc:
        raise HTTPException(status_code=422, detail=str(exc))
    finally:
        if os.path.exists(primary_path):
            os.remove(primary_path)
        if secondary_path and os.path.exists(secondary_path):
            os.remove(secondary_path)


@router.post("/validate")
async def validate_csv(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """Validate a CSV without generating a report."""
    os.makedirs(settings.UPLOAD_STORAGE_PATH, exist_ok=True)
    ext = os.path.splitext(file.filename)[1]
    tmp_path = os.path.join(
        settings.UPLOAD_STORAGE_PATH, f"validate_{uuid.uuid4().hex}{ext}"
    )
    with open(tmp_path, "wb") as f:
        shutil.copyfileobj(file.file, f)
    try:
        service = ReportService(db)
        return service.validate_csv(tmp_path)
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


@router.get("/morning-snapshots")
def list_morning_snapshots(db: Session = Depends(get_db)):
    """List all saved morning snapshots."""
    service = ReportService(db)
    return service.get_morning_snapshots()


@router.delete("/morning-snapshots/{snapshot_id}")
def delete_morning_snapshot(snapshot_id: int, db: Session = Depends(get_db)):
    """Delete a saved morning snapshot from the dashboard list."""
    service = ReportService(db)
    deleted = service.delete_morning_snapshot(snapshot_id)
    if not deleted:
        raise HTTPException(status_code=404, detail="Saved morning snapshot not found.")
    return {"message": "Saved morning snapshot deleted."}


@router.get("/history")
def get_history(skip: int = 0, limit: int = 50, db: Session = Depends(get_db)):
    reports = (
        db.query(Report)
        .order_by(Report.created_at.desc())
        .offset(skip)
        .limit(limit)
        .all()
    )
    return [
        {
            "id": r.id,
            "file_name": r.file_name,
            "sprint": r.sprint,
            "sprint_id": r.sprint_id,
            "sprint_start": r.sprint_start,
            "sprint_end": r.sprint_end,
            "day1_fixed_scope": r.day1_fixed_scope,
            "report_date": r.report_date,
            "snapshot_type": r.snapshot_type,
            "total_scope": r.total_scope,
            "completed_sp": r.completed_sp,
            "completion_percentage": r.completion_percentage,
            "created_at": r.created_at,
        }
        for r in reports
    ]


@router.get("/snapshots")
def get_snapshot_history(
    sprint_id: Optional[int] = None,
    snapshot_date: Optional[str] = None,
    snapshot_type: Optional[str] = None,
    db: Session = Depends(get_db),
):
    query = db.query(Snapshot).join(Sprint).order_by(Snapshot.captured_at.desc())
    if sprint_id is not None:
        query = query.filter(Snapshot.sprint_id == sprint_id)
    if snapshot_date:
        query = query.filter(Snapshot.snapshot_date == snapshot_date)
    if snapshot_type:
        query = query.filter(Snapshot.snapshot_type == snapshot_type.upper())
    return [_snapshot_payload(item) for item in query.all()]


@router.get("/snapshots/{snapshot_id}")
def get_snapshot(snapshot_id: int, db: Session = Depends(get_db)):
    snapshot = db.query(Snapshot).filter(Snapshot.id == snapshot_id).first()
    if not snapshot:
        raise HTTPException(status_code=404, detail="Snapshot not found.")
    return _snapshot_payload(snapshot)


@router.get("/tasks")
def get_tasks(
    snapshot_id: Optional[int] = None,
    search: Optional[str] = None,
    status: Optional[str] = None,
    issue_type: Optional[str] = None,
    developer: Optional[str] = None,
    qa_member: Optional[str] = None,
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db),
):
    """Return persisted issue snapshots for audit and task exploration."""
    limit = min(max(limit, 1), 500)
    query = db.query(IssueSnapshot).join(Snapshot).join(Sprint)
    if snapshot_id is not None:
        query = query.filter(IssueSnapshot.snapshot_id == snapshot_id)
    if search:
        term = f"%{search.strip()}%"
        query = query.filter((IssueSnapshot.issue_key.ilike(term)) | (IssueSnapshot.summary.ilike(term)))
    if status:
        query = query.filter(IssueSnapshot.status == status)
    if issue_type:
        query = query.filter(IssueSnapshot.issue_type == issue_type)
    if developer:
        query = query.filter(IssueSnapshot.developer_owner == developer)
    if qa_member:
        query = query.filter(IssueSnapshot.qa_owner == qa_member)
    total = query.count()
    rows = query.order_by(IssueSnapshot.issue_key.asc()).offset(skip).limit(limit).all()
    return {
        "total": total,
        "skip": skip,
        "limit": limit,
        "tasks": [
            {
                "id": row.id,
                "snapshot_id": row.snapshot_id,
                "issue_key": row.issue_key,
                "summary": row.summary,
                "issue_type": row.issue_type,
                "parent_key": row.parent_key,
                "status": row.status,
                "assignee": row.assignee,
                "developer_owner": row.developer_owner,
                "developer_sp": row.developer_sp,
                "qa_owner": row.qa_owner,
                "qa_sp": row.qa_sp,
                "story_points": row.story_points,
                "completed_sp": row.completed_sp,
                "remaining_sp": max(0.0, (row.story_points or 0.0) - (row.completed_sp or 0.0)),
                "sprint": row.snapshot.sprint.name,
                "snapshot_date": row.snapshot.snapshot_date,
                "snapshot_type": row.snapshot.snapshot_type,
            }
            for row in rows
        ],
    }


@router.get("/sprints")
def get_sprint_history(db: Session = Depends(get_db)):
    sprints = db.query(Sprint).order_by(Sprint.created_at.desc()).all()
    result = []
    for sprint in sprints:
        snapshots = sorted(sprint.snapshots, key=lambda item: item.captured_at or datetime.min, reverse=True)
        latest = snapshots[0] if snapshots else None
        result.append({
            "id": sprint.id,
            "sprint_key": sprint.sprint_key,
            "name": sprint.name,
            "sprint_id": sprint.sprint_id,
            "start_date": sprint.start_date,
            "end_date": sprint.end_date,
            "day1_fixed_scope": sprint.day1_fixed_scope,
            "snapshot_count": len(snapshots),
            "current_scope": latest.overall_scope if latest else 0.0,
            "completed_sp": latest.overall_completed if latest else 0.0,
            "remaining_sp": latest.overall_remaining if latest else 0.0,
            "completion_pct": latest.overall_completion if latest else 0.0,
            "last_updated": latest.captured_at if latest else sprint.created_at,
        })
    return result


@router.get("/{report_id}")
def get_report(report_id: int, db: Session = Depends(get_db)):
    report = db.query(Report).filter(Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found.")
    return report


@router.get("/{report_id}/download")
def download_report(report_id: int, db: Session = Depends(get_db)):
    report = db.query(Report).filter(Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found.")
    if not report.file_path or not os.path.exists(report.file_path):
        raise HTTPException(status_code=404, detail="Excel file not found on disk.")
    return FileResponse(
        path=report.file_path,
        filename=report.file_name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@router.get("/{report_id}/preview", response_class=HTMLResponse)
def preview_report(report_id: int, db: Session = Depends(get_db)):
    """Render the generated Excel report as a browser-viewable HTML preview."""
    report = db.query(Report).filter(Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found.")
    if not report.file_path or not os.path.exists(report.file_path):
        raise HTTPException(status_code=404, detail="Excel file not found on disk.")

    from openpyxl import load_workbook

    # Keep formula source available for the download, but render the simple
    # report formulas as values in the browser preview.
    workbook = load_workbook(report.file_path, read_only=True, data_only=False)

    def display_value(worksheet, cell):
        value = cell.value
        if not (isinstance(value, str) and value.startswith("=")):
            return value

        expression = value[1:].strip()
        if expression.startswith("IFERROR(") and expression.endswith(',"N/A")'):
            expression = expression[len("IFERROR("):-len(',"N/A")')]
        evaluated = evaluate_expression(worksheet, expression)
        if (
            isinstance(evaluated, (int, float))
            and "%" in (cell.number_format or "")
        ):
            return f"{evaluated * 100:.2f}%"
        return evaluated

    def evaluate_expression(worksheet, expression):
        if "/" in expression and not expression.startswith("SUM("):
            try:
                left, right = expression.split("/", 1)
                numerator = float(display_value(worksheet, worksheet[left]) or 0)
                denominator = float(display_value(worksheet, worksheet[right]) or 0)
                return numerator / denominator if denominator else "N/A"
            except (ValueError, ZeroDivisionError):
                return "N/A"
        if expression.startswith("SUM(") and expression.endswith(")"):
            try:
                total = 0.0
                for row in worksheet[expression[4:-1]]:
                    values = row if isinstance(row, tuple) else (row,)
                    total += sum(float(display_value(worksheet, item) or 0) for item in values)
                return total
            except (TypeError, ValueError, KeyError):
                return "N/A"
        if "-" in expression:
            try:
                left, right = expression.split("-", 1)
                return float(display_value(worksheet, worksheet[left]) or 0) - float(
                    display_value(worksheet, worksheet[right]) or 0
                )
            except (TypeError, ValueError, KeyError):
                return "N/A"
        return value

    tables = []
    for worksheet in workbook.worksheets:
        rows = []
        for row in worksheet.iter_rows():
            cells = "".join(
                f"<td>{escape('' if display_value(worksheet, cell) is None else str(display_value(worksheet, cell)))}</td>"
                for cell in row
            )
            if cells.replace("<td></td>", ""):
                rows.append(f"<tr>{cells}</tr>")
        tables.append(
            f"<section><h2>{escape(worksheet.title)}</h2>"
            f"<div class='table-wrap'><table>{''.join(rows)}</table></div></section>"
        )
    workbook.close()
    title = escape(report.file_name or "Sprint Report")
    return HTMLResponse(
        f"""<!doctype html><html><head><meta charset="utf-8"><title>{title}</title>
<style>
body{{font-family:Arial,sans-serif;background:#f3f6fa;color:#172033;margin:0;padding:24px}}
main{{max-width:1400px;margin:auto;background:#fff;padding:24px;border-radius:12px;box-shadow:0 2px 12px #0001}}
h1{{margin-top:0;color:#1f3864}} h2{{color:#2e75b6;margin-top:28px}}
.table-wrap{{overflow:auto;border:1px solid #d8dee8;border-radius:8px}}
table{{border-collapse:collapse;min-width:900px;width:100%}} td{{border:1px solid #d8dee8;padding:8px;white-space:pre-wrap}}
tr:first-child td{{font-weight:700;background:#1f3864;color:#fff}}
</style></head><body><main><h1>{title}</h1>{''.join(tables)}
<p><a href="/api/reports/{report_id}/download">Download original Excel file</a></p>
</main></body></html>"""
    )


@router.delete("/{report_id}")
def delete_report(report_id: int, db: Session = Depends(get_db)):
    report = db.query(Report).filter(Report.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Report not found.")

    file_path = report.file_path
    if file_path:
        try:
            _delete_file_with_retry(file_path)
        except FileNotFoundError:
            pass
        except PermissionError as exc:
            raise HTTPException(
                status_code=409,
                detail="The report file is currently in use and cannot be deleted. Please close the file and try again.",
            ) from exc

    db.delete(report)
    db.commit()
    return {"message": "Report deleted."}
